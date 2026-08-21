from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\ptbxl")
SRC = ROOT / "results" / "ACTIVITY_REPORT.md"
OUT = ROOT / "results" / "ACTIVITY_REPORT.docx"

BLUE = "2E74B5"; DARK = "1F4D78"; LIGHT = "E8EEF5"; GRAY = "F2F4F7"


def rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi"); ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def font_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Tahoma"; run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Tahoma")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); shd = tcpr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tcpr=cell._tc.get_or_add_tcPr(); tcw=tcpr.find(qn("w:tcW"))
    if tcw is None: tcw=OxmlElement("w:tcW"); tcpr.append(tcw)
    tcw.set(qn("w:w"),str(dxa)); tcw.set(qn("w:type"),"dxa")


def set_table_geometry(table, widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    tblpr=table._tbl.tblPr; tblw=tblpr.find(qn("w:tblW"))
    if tblw is None: tblw=OxmlElement("w:tblW"); tblpr.append(tblw)
    tblw.set(qn("w:w"),str(sum(widths))); tblw.set(qn("w:type"),"dxa")
    ind=tblpr.find(qn("w:tblInd"))
    if ind is None: ind=OxmlElement("w:tblInd"); tblpr.append(ind)
    ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa")
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col=OxmlElement("w:gridCol"); col.set(qn("w:w"),str(w)); grid.append(col)
    for row in table.rows:
        for i,cell in enumerate(row.cells): set_cell_width(cell,widths[i])


def add_page_field(paragraph):
    run=paragraph.add_run(); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); run._r.append(fld)


def clean_md(s):
    return re.sub(r"\*\*(.*?)\*\*", r"\1", s).replace("`", "")


doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)

styles=doc.styles
normal=styles["Normal"]; normal.font.name="Tahoma"; normal._element.rPr.rFonts.set(qn("w:cs"),"Tahoma"); normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in [("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,DARK,10,5)]:
    st=styles[name]; st.font.name="Tahoma"; st._element.rPr.rFonts.set(qn("w:cs"),"Tahoma"); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

# Running header/footer.
hp=sec.header.paragraphs[0]; rtl(hp); font_run(hp.add_run("گزارش فعالیت پژوهشی | پروژه PTB-XL"),9,False,"666666")
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font_run(fp.add_run("صفحه "),9,False,"666666"); add_page_field(fp)

# Editorial cover.
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.CENTER); font_run(p.add_run("گزارش جامع فعالیت‌های پروژه PTB-XL"),25,True,DARK)
p.paragraph_format.space_after=Pt(12)
p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.CENTER); font_run(p.add_run("توسعه Teacher و آماده‌سازی Multi-Level Knowledge Distillation"),14,False,BLUE)
p.paragraph_format.space_after=Pt(24)
p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.CENTER); font_run(p.add_run("گزارش فنی اجرا، خطاها، تصمیم‌ها و نتایج اعتبارسنجی"),11,False,"555555")
p.paragraph_format.space_after=Pt(70)
p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.CENTER); font_run(p.add_run("نسخه ثبت‌شده: ۲۱ اوت ۲۰۲۶\nوضعیت Test: ارزیابی نشده"),11,True,DARK)
doc.add_page_break()

lines=SRC.read_text(encoding="utf-8").splitlines()[4:]
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith("|"):
        block=[]
        while i<len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip()); i+=1
        rows=[]
        for j,b in enumerate(block):
            cells=[clean_md(x.strip()) for x in b.strip("|").split("|")]
            if j==1 and all(set(x)<=set("-:") for x in cells): continue
            rows.append(cells)
        cols=len(rows[0]); table=doc.add_table(rows=len(rows),cols=cols); table.style="Table Grid"
        if cols==2: widths=[2700,6660]
        elif cols==3: widths=[2400,3480,3480]
        elif cols==4: widths=[1900,2486,2487,2487]
        elif cols==5: widths=[1700,1915,1915,1915,1915]
        elif cols==6: widths=[1300,1612,1612,1612,1612,1612]
        else: widths=[9360//cols]*cols; widths[-1]+=9360-sum(widths)
        for r,row in enumerate(rows):
            for c,text in enumerate(row):
                cell=table.cell(r,c); cell.text=""; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0]; rtl(p); font_run(p.add_run(text),8.7,r==0,DARK if r==0 else "000000")
                if r==0: set_cell_shading(cell,LIGHT)
                elif r%2==0: set_cell_shading(cell,"F8FAFC")
        set_table_geometry(table,widths)
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
        continue
    if line.startswith("### "):
        p=doc.add_paragraph(style="Heading 3"); rtl(p); font_run(p.add_run(clean_md(line[4:])),12,True,DARK)
    elif line.startswith("## "):
        p=doc.add_paragraph(style="Heading 2"); rtl(p); font_run(p.add_run(clean_md(line[3:])),13,True,BLUE)
    elif line.startswith("# "):
        p=doc.add_paragraph(style="Heading 1"); rtl(p); font_run(p.add_run(clean_md(line[2:])),16,True,BLUE)
    elif re.match(r"^\d+\. ",line):
        p=doc.add_paragraph(style="List Number"); rtl(p); font_run(p.add_run(clean_md(re.sub(r"^\d+\. ","",line))),10.5)
    elif line.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); rtl(p); font_run(p.add_run(clean_md(line[2:])),10.5)
    else:
        p=doc.add_paragraph(); rtl(p); font_run(p.add_run(clean_md(line)),10.5)
    i+=1

props=doc.core_properties; props.title="گزارش جامع فعالیت‌های پروژه PTB-XL"; props.subject="Teacher training and knowledge distillation activity report"; props.author="PTB-XL Research Project"
doc.save(OUT)
print(OUT)
